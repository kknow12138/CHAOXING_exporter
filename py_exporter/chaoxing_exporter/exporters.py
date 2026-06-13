"""把解析好的题目导出为 Word(.docx) / JSON / Markdown。"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Iterable, List

from .parser import Question


def _safe_filename(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|]+', "_", name).strip()
    return name or "作业"


def export_json(questions: List[Question], out_path: str | Path, *, title: str = "") -> Path:
    out_path = Path(out_path)
    payload = {
        "title": title,
        "count": len(questions),
        "questions": [q.to_dict() for q in questions],
    }
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return out_path


def export_markdown(questions: List[Question], out_path: str | Path, *, title: str = "") -> Path:
    out_path = Path(out_path)
    lines: List[str] = []
    if title:
        lines.append(f"# {title}\n")
    for q in questions:
        lines.append(f"## {q.number}. [{q.type}] {q.title}")
        if q.options:
            for opt in q.options:
                lines.append(f"- {opt}")
        if q.answer:
            lines.append(f"\n**正确答案：** {q.answer}")
        if q.my_answer:
            lines.append(f"**我的答案：** {q.my_answer}")
        lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path


def export_word(questions: List[Question], out_path: str | Path, *, title: str = "") -> Path:
    from docx import Document
    from docx.shared import Pt

    out_path = Path(out_path)
    doc = Document()
    if title:
        doc.add_heading(title, level=0)

    for q in questions:
        h = doc.add_paragraph()
        run = h.add_run(f"{q.number}. [{q.type}] {q.title}")
        run.bold = True
        run.font.size = Pt(12)

        for opt in q.options:
            doc.add_paragraph(opt, style="List Bullet")

        if q.answer:
            p = doc.add_paragraph()
            p.add_run("正确答案：").bold = True
            p.add_run(q.answer)
        if q.my_answer:
            p = doc.add_paragraph()
            p.add_run("我的答案：").bold = True
            p.add_run(q.my_answer)
        doc.add_paragraph("")

    doc.save(str(out_path))
    return out_path


_EXPORTERS = {
    "word": (export_word, ".docx"),
    "json": (export_json, ".json"),
    "markdown": (export_markdown, ".md"),
}


def export(
    questions: List[Question],
    out_dir: str | Path,
    base_name: str,
    formats: Iterable[str],
    *,
    title: str = "",
) -> List[Path]:
    """按指定格式导出，返回生成的文件路径列表。"""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = _safe_filename(base_name)
    written: List[Path] = []
    for fmt in formats:
        if fmt not in _EXPORTERS:
            continue
        func, ext = _EXPORTERS[fmt]
        written.append(func(questions, out_dir / f"{base}{ext}", title=title or base_name))
    return written
